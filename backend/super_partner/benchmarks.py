"""Public benchmark results, kept separate from the application's routing presets."""

from __future__ import annotations

import asyncio
import csv
import hashlib
import io
import json
import math
import os
import re
import unicodedata
import zipfile
from datetime import date, datetime
from html.parser import HTMLParser
from typing import Literal, Optional
from urllib.parse import quote

import httpx
from openpyxl import load_workbook
from pydantic import BaseModel, Field

from .config import settings
from .db import Database
from .models import ModelProfile, utcnow

BFCL_URL = "https://gorilla.cs.berkeley.edu/leaderboard.html"
BFCL_DATA_URL = "https://gorilla.cs.berkeley.edu/data_overall.csv"
SUPERCLUE_URL = "https://www.superclueai.com/"
AA_URL = "https://artificialanalysis.ai/api/v2/language/models/free"
AA_PUBLIC_URL = "https://artificialanalysis.ai/leaderboards/models"
AA_KEY_ENV = "ARTIFICIAL_ANALYSIS_API_KEY"
SOURCE_NAMES = {
    "superclue": "SuperCLUE",
    "bfcl": "BFCL",
    "artificial_analysis": "Artificial Analysis",
}


def aa_key() -> Optional[str]:
    return os.getenv(AA_KEY_ENV) or (
        settings.artificial_analysis_api_key.get_secret_value()
        if settings.artificial_analysis_api_key
        else None
    )


class BenchmarkMetric(BaseModel):
    key: str
    label: str
    value: float = Field(allow_inf_nan=False)
    unit: Literal["percent", "index"]
    note: Optional[str] = None
    benchmark: Optional[str] = None


class BenchmarkResult(BaseModel):
    id: str
    source: str
    source_name: str
    source_url: str
    data_url: str
    version: str
    model_name: str
    provider: Optional[str] = None
    model_key: str
    variant: str = "来源未注明配置"
    published_at: Optional[date] = None
    fetched_at: datetime = Field(default_factory=utcnow)
    metrics: list[BenchmarkMetric]
    collection_method: Literal["dataset", "api", "public_page"] = "dataset"


# A leaderboard name ending in a snapshot date: "DeepSeek V4 Pro 0813", "... 20250813".
SNAPSHOT = re.compile(r"(.*?)[\s_-]+(\d{4}|\d{8})")


def normalized(value: str) -> str:
    # Formatting differences only: never drop snapshot dates, model tiers, or version digits.
    return re.sub(r"[\s._-]+", "", unicodedata.normalize("NFKC", value).casefold())


PROVIDERS = {
    "openai": "openai",
    "anthropic": "anthropic",
    "google": "google",
    "googledeepmind": "google",
    "deepseek": "deepseek",
    "深度求索": "deepseek",
    "alibaba": "qwen",
    "alibabacloud": "qwen",
    "阿里巴巴": "qwen",
    "阿里云": "qwen",
    "字节跳动": "doubao",
    "bytedance": "doubao",
    "火山引擎": "doubao",
    "zai": "glm",
    "zhipuai": "glm",
    "智谱ai": "glm",
    "智谱": "glm",
    "minimax": "minimax",
    "稀宇科技": "minimax",
    "moonshotai": "kimi",
    "月之暗面": "kimi",
    "kimi": "kimi",
}


def parse_identity(name: str) -> tuple[str, str]:
    name = unicodedata.normalize("NFKC", name).strip()
    # Only known inference settings can be removed; dates/editions in parentheses are identity.
    match = re.fullmatch(r"(.+?)\s*\(([^()]+)\)", name)
    if match and re.fullmatch(
        r"FC(?: thinking)?|Prompt(?: \+ Thinking)?|"
        r"(?:minimal|low|medium|high|xhigh|max)(?: with fallback)?|"
        r"non[- ]reasoning|reasoning|thinking",
        match[2],
        re.IGNORECASE,
    ):
        return match[1].strip(), match[2].strip()
    return name, "来源未注明配置"


def metric_value(value, unit: str) -> Optional[float]:
    if value is None or isinstance(value, bool) or str(value).strip() in {"", "N/A", "—", "-"}:
        return None
    result = float(str(value).strip().rstrip("%"))
    if not math.isfinite(result) or (unit == "percent" and not 0 <= result <= 100):
        raise ValueError("Invalid benchmark score")
    return result


def metrics(row: dict, fields: dict[str, str], unit: str = "percent") -> list[BenchmarkMetric]:
    result = []
    for key, label in fields.items():
        value = metric_value(row.get(key), unit)
        if value is not None:
            result.append(BenchmarkMetric(key=key, label=label, value=value, unit=unit))
    return result


def result_for(
    source: str, name: str, organization: str, version: str, **kwargs
) -> BenchmarkResult:
    key, variant = parse_identity(name)
    identifier = hashlib.sha256(f"{source}|{version}|{name}".encode()).hexdigest()[:20]
    return BenchmarkResult(
        id=identifier,
        source=source,
        source_name=SOURCE_NAMES[source],
        version=version,
        model_name=name,
        model_key=key,
        variant=variant,
        provider=PROVIDERS.get(normalized(organization)),
        **kwargs,
    )


BFCL_FIELDS = {
    "Overall Acc": "工具调用综合",
    "Non-Live AST Acc": "单轮标准任务",
    "Live Acc": "单轮真实任务",
    "Multi Turn Acc": "多轮工具调用",
    "Web Search Acc": "网页搜索任务",
    "Memory Acc": "记忆任务",
}
SUPERCLUE_FIELDS = {
    key: key
    for key in (
        "总分",
        "数学推理",
        "幻觉控制",
        "科学推理",
        "精确指令遵循",
        "智能体编程",
        "智能体任务规划",
    )
}
AA_FIELDS = {
    "artificial_analysis_intelligence_index": "综合智能指数",
    "artificial_analysis_coding_index": "编程指数",
    "artificial_analysis_agentic_index": "智能体指数",
}
AA_DETAIL_FIELDS = {
    "scicode": ("aa_scicode", "科学编程", "SciCode"),
    "terminalBench40": ("aa_terminalbench_4", "终端智能体", "Terminal-Bench 4.0"),
    "hle": ("aa_hle", "复杂推理", "HLE"),
    "critpt": ("aa_critpt", "物理推理", "CritPt"),
    "lcr": ("aa_lcr", "长文本推理", "AA-LCR"),
    "omniscienceAccuracy": ("aa_knowledge_accuracy", "知识准确率", "AA-Omniscience Accuracy"),
    "omniscienceNonHallucination": (
        "aa_non_hallucination",
        "非幻觉率",
        "AA-Omniscience Non-Hallucination",
    ),
}


def parse_bfcl(text: str, page: str) -> list[BenchmarkResult]:
    reader = csv.DictReader(io.StringIO(text.lstrip("\ufeff")))
    if not {"Model", "Organization", *BFCL_FIELDS}.issubset(reader.fieldnames or []):
        raise ValueError("BFCL schema changed")
    version = re.search(r"BFCL\)?\s*(V\d+)", page)
    stamp = re.search(r"Last Updated:\s*(\d{4}-\d{2}-\d{2})", page)
    if not version or not stamp:
        raise ValueError("Missing BFCL version/date")
    result = []
    for row in reader:
        name = row.get("Model", "").strip()
        if not name:
            continue
        scores = metrics(row, BFCL_FIELDS)
        if scores:
            result.append(
                result_for(
                    "bfcl",
                    name,
                    row.get("Organization", ""),
                    f"BFCL {version[1]}",
                    source_url=BFCL_URL,
                    data_url=BFCL_DATA_URL,
                    published_at=date.fromisoformat(stamp[1]),
                    metrics=scores,
                )
            )
    if not result:
        raise ValueError("Empty BFCL results")
    return result


def parse_superclue(content: bytes, edition: str, data_url: str) -> list[BenchmarkResult]:
    # Bound workbook expansion; read cached values without executing formulas or external links.
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 32_000_000:
            raise ValueError("Workbook exceeds size limit")
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        if "总排行榜" not in workbook.sheetnames:
            raise ValueError("SuperCLUE schema changed")
        rows = workbook["总排行榜"].iter_rows(values_only=True)
        headers = [str(cell or "").strip() for cell in next(rows)]
        if not {"模型名称", "机构", *SUPERCLUE_FIELDS}.issubset(headers):
            raise ValueError("SuperCLUE columns changed")
        result = []
        for index, values in enumerate(rows):
            if index > 10000:
                raise ValueError("Too many benchmark rows")
            row = dict(zip(headers, values))
            if not row.get("模型名称"):
                continue
            scores = metrics(row, SUPERCLUE_FIELDS)
            stamp = row.get("发布日期") or row.get("测评发布日期")
            if isinstance(stamp, datetime):
                published = stamp.date()
            elif isinstance(stamp, date):
                published = stamp
            elif stamp:
                parts = re.fullmatch(r"(\d{4})[./-](\d{1,2})[./-](\d{1,2})", str(stamp).strip())
                if not parts:
                    raise ValueError("Invalid evaluation date")
                published = date(*(int(part) for part in parts.groups()))
            else:
                published = None
            if scores:
                result.append(
                    result_for(
                        "superclue",
                        str(row["模型名称"]),
                        str(row["机构"]),
                        f"SuperCLUE {edition}",
                        source_url=SUPERCLUE_URL,
                        data_url=data_url,
                        published_at=published,
                        metrics=scores,
                    )
                )
        if not result:
            raise ValueError("Empty SuperCLUE results")
        return result
    finally:
        workbook.close()


def parse_aa(payload: dict) -> list[BenchmarkResult]:
    if (
        not isinstance(payload.get("data"), list)
        or payload.get("intelligence_index_version") is None
    ):
        raise ValueError("AA schema changed")
    version = f"AA Intelligence Index v{payload['intelligence_index_version']}"
    result = []
    for row in payload["data"]:
        scores = metrics(row["evaluations"], AA_FIELDS, "index")
        if scores:
            result.append(
                result_for(
                    "artificial_analysis",
                    row["name"],
                    row["model_creator"]["name"],
                    version,
                    source_url=f"https://artificialanalysis.ai/models/{quote(row['slug'], safe='')}",
                    data_url=AA_URL,
                    metrics=scores,
                    collection_method="api",
                )
            )
    return result


class PublicTableParser(HTMLParser):
    """Read public HTML and serialized chart data without executing JavaScript."""

    def __init__(self):
        super().__init__()
        self.tables = []
        self.table = None
        self.row = None
        self.cell = None
        self.scripts = []
        self.script = None

    def handle_starttag(self, tag, attrs):
        if tag == "script":
            self.script = []
        elif tag == "table":
            if self.table is not None:
                raise ValueError("Nested leaderboard table")
            self.table = []
        elif self.table is not None:
            if tag == "tr":
                self.row = []
            elif tag in {"td", "th"}:
                self.cell = {"text": [], "links": []}
            elif tag == "a" and self.cell is not None:
                self.cell["links"].append(dict(attrs).get("href", ""))

    def handle_data(self, data):
        if self.script is not None:
            self.script.append(data)
        elif self.cell is not None:
            self.cell["text"].append(data)

    def handle_endtag(self, tag):
        if tag == "script" and self.script is not None:
            self.scripts.append("".join(self.script))
            self.script = None
        elif tag in {"td", "th"} and self.cell is not None:
            if self.row is None:
                raise ValueError("Leaderboard cell outside row")
            self.cell["text"] = " ".join(" ".join(self.cell["text"]).split())
            self.row.append(self.cell)
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None


def aa_public_details(scripts: list[str]) -> list[dict]:
    """Decode only JSON in the public page's Next.js stream, never eval script text."""
    fragments = []
    for script in scripts:
        match = re.fullmatch(r"self\.__next_f\.push\((\[.*\])\);?", script.strip(), re.DOTALL)
        if not match:
            continue
        try:
            frame = json.loads(match[1])
        except ValueError:
            continue
        if isinstance(frame, list) and len(frame) == 2 and frame[0] == 1 and isinstance(frame[1], str):
            fragments.append(frame[1])
    rows = []
    for line in "".join(fragments).splitlines():
        record = re.fullmatch(r"[0-9a-f]+:(.*)", line)
        if not record:
            continue
        try:
            pending = [json.loads(record[1])]
        except ValueError:
            continue  # Other stream records include module identifiers and plain text.
        while pending:
            value = pending.pop()
            if isinstance(value, dict):
                if {"slug", "shortName", "modelCreatorName", "intelligenceIndex"} <= value.keys():
                    rows.append(value)
                else:
                    pending.extend(value.values())
            elif isinstance(value, list):
                pending.extend(value)
    return rows


def add_aa_details(results: list[BenchmarkResult], details: list[dict]) -> list[BenchmarkResult]:
    results = [result.model_copy(deep=True) for result in results]
    by_slug = {}
    for row in details:
        if not all(
            isinstance(row.get(key), str) for key in ("slug", "shortName", "modelCreatorName")
        ):
            continue
        by_slug.setdefault(row["slug"], []).append(row)
    for result in results:
        candidates = by_slug.get(result.source_url.rsplit("/", 1)[-1], [])
        if len(candidates) != 1:
            continue
        row = candidates[0]
        if PROVIDERS.get(normalized(row["modelCreatorName"])) != result.provider or normalized(
            result.model_name
        ) not in {normalized(row["shortName"]), normalized(row.get("name") or "")}:
            continue
        index = row["intelligenceIndex"]
        headline = next(
            (m.value for m in result.metrics if m.key == "artificial_analysis_intelligence_index"),
            None,
        )
        if (
            headline is None
            or isinstance(index, bool)
            or not isinstance(index, (int, float))
            or not math.isfinite(index)
            or abs(index - headline) > 0.500001
        ):
            continue  # Reject mismatched snapshots; the public table rounds the index.
        for field, (key, label, benchmark) in AA_DETAIL_FIELDS.items():
            value = row.get(field)
            if value is None:
                continue
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or not 0 <= value <= 1
            ):
                raise ValueError("Invalid AA benchmark ratio")
            result.metrics.append(
                BenchmarkMetric(
                    key=key,
                    label=label,
                    value=value * 100,
                    unit="percent",
                    benchmark=benchmark,
                )
            )
    return results


def parse_aa_public(page: str) -> list[BenchmarkResult]:
    parser = PublicTableParser()
    parser.feed(page)
    parser.close()
    if parser.table is not None:
        raise ValueError("Incomplete AA leaderboard")
    required = {"Model", "Creator", "Artificial Analysis Intelligence Index", "Further Analysis"}
    candidates = []
    for table in parser.tables:
        for index, row in enumerate(table):
            headers = [cell["text"] for cell in row]
            if required.issubset(headers):
                candidates.append((headers, table[index + 1 :]))
    if len(candidates) != 1:
        raise ValueError("AA public leaderboard columns changed")
    headers, rows = candidates[0]
    results = []
    for cells in rows:
        if len(cells) != len(headers):
            raise ValueError("Incomplete AA leaderboard row")
        row = dict(zip(headers, cells))
        name, organization = row["Model"]["text"], row["Creator"]["text"]
        raw = row["Artificial Analysis Intelligence Index"]["text"]
        if raw in {"", "--", "—", "N/A"}:
            continue  # Listed without a score is not a zero or an evaluated model.
        number = re.fullmatch(r"(\d+(?:\.\d+)?)\s*(\*)?", raw)
        links = [
            link
            for link in row["Further Analysis"]["links"]
            if re.fullmatch(r"/models/[a-z0-9_-]+", link)
        ]
        if not number or not name or not organization or len(links) != 1:
            raise ValueError("Invalid AA public leaderboard row")
        source_url = "https://artificialanalysis.ai" + links[0]
        result = result_for(
            "artificial_analysis",
            name,
            organization,
            "AA Intelligence Index（网页未标版本）",
            source_url=source_url,
            data_url=AA_PUBLIC_URL,
            collection_method="public_page",
            metrics=[
                BenchmarkMetric(
                    key="artificial_analysis_intelligence_index",
                    label="综合智能指数",
                    value=float(number[1]),
                    unit="index",
                    note="来源带 * 标记，请参阅原榜单说明" if number[2] else None,
                )
            ],
        )
        # Distinct source model pages can share a display label; keep their identities separate.
        result.id = hashlib.sha256(f"{result.id}|{source_url}".encode()).hexdigest()[:20]
        results.append(result)
    if not results:
        raise ValueError("Empty AA public leaderboard")
    identities = {}
    for result in results:
        identities.setdefault((result.provider, result.model_name), set()).add(result.source_url)
    for result in results:
        if len(identities[(result.provider, result.model_name)]) > 1:
            result.model_key = result.source_url.rsplit("/", 1)[-1]
    return add_aa_details(results, aa_public_details(parser.scripts))


def match_results(
    models: list[ModelProfile], results: list[BenchmarkResult]
) -> dict[str, list[dict]]:
    ids, names = {}, {}
    by_id = {model.id: model for model in models}
    for model in models:
        if model.channel != "api" or model.provider == "mock":
            continue
        ids.setdefault((model.provider, normalized(model.id.split(":", 1)[1])), set()).add(model.id)
        names.setdefault((model.provider, normalized(model.display_name)), set()).add(model.id)
    matched = {}
    for result in results:
        key = (result.provider, normalized(result.model_key))
        candidates = ids.get(key) or names.get(key) or set()
        # Ambiguous display names do not justify copying scores to multiple model snapshots.
        if len(candidates) == 1:
            identifier = next(iter(candidates))
            model = by_id[identifier]
            if (
                model.release_date
                and result.published_at
                and result.published_at < model.release_date
            ):
                continue
            matched.setdefault(identifier, []).append(
                {
                    **result.model_dump(mode="json"),
                    "matched_by": "model_id" if ids.get(key) else "display_name",
                }
            )
    # An undated alias (deepseek-v4-pro) serves the provider's newest snapshot, which a
    # leaderboard may list with its date ("DeepSeek V4 Pro 0813"). Unless the catalog lists
    # that newest snapshot itself, the alias takes its results.
    taken = {row["id"] for rows in matched.values() for row in rows}
    latest = {}
    for result in results:
        dated = SNAPSHOT.fullmatch(result.model_key.strip())
        if not dated:
            continue
        key = (result.provider, normalized(dated[1]))
        candidates = ids.get(key) or names.get(key) or set()
        identifier = next(iter(candidates)) if len(candidates) == 1 else None
        # An exact match from the same source wins over a snapshot.
        if identifier is None or any(row["source"] == result.source for row in matched.get(identifier, [])):
            continue
        newest = latest.get((identifier, result.source))
        if newest is None or int(dated[2]) > newest[0]:
            newest = latest[identifier, result.source] = (int(dated[2]), [])
        if int(dated[2]) == newest[0]:
            newest[1].append(result)
    for (identifier, _source), (_date, rows) in latest.items():
        if any(result.id in taken for result in rows):
            continue
        matched.setdefault(identifier, []).extend(
            {**result.model_dump(mode="json"), "matched_by": "latest_snapshot"} for result in rows
        )
    for rows in matched.values():
        rows.sort(key=lambda row: (row["source"], row["version"], row["model_name"]))
    return matched


class BenchmarkCatalog:
    def __init__(self, db: Database):
        self.db = db
        self.cache = db.load_benchmark_cache()

    def results(self) -> list[BenchmarkResult]:
        return [
            BenchmarkResult.model_validate(row)
            for entry in self.cache.values()
            for row in entry.get("results", [])
        ]

    def match_models(self, models: list[ModelProfile]) -> dict[str, list[dict]]:
        return match_results(models, self.results())

    def status(self, models: list[ModelProfile]) -> dict:
        matched = self.match_models(models)
        return {
            "matched_models": len(matched),
            "sources": {
                source: {
                    **{
                        key: value
                        for key, value in self.cache.get(source, {}).items()
                        if key != "results"
                    },
                    "name": name,
                    "matched_models": sum(
                        any(r["source"] == source for r in rows) for rows in matched.values()
                    ),
                    "configured": bool(aa_key()) if source == "artificial_analysis" else True,
                }
                for source, name in SOURCE_NAMES.items()
            },
        }

    async def _get(self, client: httpx.AsyncClient, url: str, **kwargs) -> httpx.Response:
        async with client.stream("GET", url, follow_redirects=False, **kwargs) as response:
            response.raise_for_status()
            content = bytearray()
            async for chunk in response.aiter_bytes():
                content.extend(chunk)
                if len(content) > 8_000_000:
                    raise ValueError("Benchmark response exceeds size limit")
            # aiter_bytes already decompresses the body; don't ask Response to decode it again.
            return httpx.Response(
                200,
                content=bytes(content),
                headers={
                    "content-type": response.headers.get("content-type", "application/octet-stream")
                },
            )

    async def _superclue(self, client: httpx.AsyncClient) -> list[BenchmarkResult]:
        page = (await self._get(client, SUPERCLUE_URL)).text
        entry = re.search(r'src="(/assets/index-[\w-]+\.js)"', page)
        if not entry:
            raise ValueError("SuperCLUE entry script changed")
        entry_text = (await self._get(client, SUPERCLUE_URL.rstrip("/") + entry[1])).text
        home = re.search(r"\./(HomePage-[\w-]+\.js)", entry_text)
        if not home:
            raise ValueError("SuperCLUE home script changed")
        home_text = (await self._get(client, SUPERCLUE_URL + "assets/" + home[1])).text
        chart = re.search(r"\./(charts-main-[\w-]+\.js)", home_text)
        if not chart:
            raise ValueError("SuperCLUE chart script changed")
        chart_text = (await self._get(client, SUPERCLUE_URL + "assets/" + chart[1])).text
        editions = re.findall(r'loadExcelData\("(\d{4}年\d{1,2}月)"', chart_text)
        if len(set(editions)) != 1:
            raise ValueError("Ambiguous current SuperCLUE edition")
        edition = editions[0]
        url = SUPERCLUE_URL + "data/generalboard/" + quote(edition) + ".xlsx"
        content = (await self._get(client, url)).content
        return await asyncio.to_thread(parse_superclue, content, edition, url)

    async def _bfcl(self, client: httpx.AsyncClient) -> list[BenchmarkResult]:
        page, data = await asyncio.gather(
            self._get(client, BFCL_URL), self._get(client, BFCL_DATA_URL)
        )
        return parse_bfcl(data.text, page.text)

    async def _aa(self, client: httpx.AsyncClient) -> list[BenchmarkResult]:
        if not aa_key():
            page = (await self._get(client, AA_PUBLIC_URL)).text
            return await asyncio.to_thread(parse_aa_public, page)
        results, version = [], None
        for page in range(1, 21):
            response = await self._get(
                client, AA_URL, headers={"x-api-key": aa_key()}, params={"page": page}
            )
            body = response.json()
            if version is not None and body.get("intelligence_index_version") != version:
                raise ValueError("Benchmark version changed during pagination")
            version = body.get("intelligence_index_version")
            results.extend(parse_aa(body))
            pagination = body["pagination"]
            if pagination["page"] != page or not isinstance(pagination["has_more"], bool):
                raise ValueError("Invalid pagination")
            if not pagination["has_more"]:
                # Public per-benchmark measurements also complement the API's headline indices.
                try:
                    page = (await self._get(client, AA_PUBLIC_URL)).text
                    parser = PublicTableParser()
                    parser.feed(page)
                    parser.close()
                    return add_aa_details(results, aa_public_details(parser.scripts))
                except (httpx.HTTPError, ValueError):
                    return results
        raise ValueError("Too many benchmark pages")

    async def refresh(self, client: httpx.AsyncClient) -> None:
        async def sync(source, fetch):
            previous = self.cache.get(source, {})
            entry = {**previous, "checked_at": utcnow().isoformat()}
            try:
                rows = await asyncio.wait_for(fetch(client), timeout=60)
                if not rows or len({row.id for row in rows}) != len(rows):
                    raise ValueError("Empty or duplicate results")
                published = [row.published_at.isoformat() for row in rows if row.published_at]
                entry.update(
                    status="ok",
                    error=None,
                    last_success_at=utcnow().isoformat(),
                    results=[row.model_dump(mode="json") for row in rows],
                    count=len(rows),
                    with_details=sum(any(m.benchmark for m in row.metrics) for row in rows),
                    collection_method=rows[0].collection_method,
                    latest_published_at=max(published, default=None),
                )
            except Exception as exc:
                if isinstance(exc, httpx.HTTPStatusError):
                    message = f"HTTP {exc.response.status_code}"
                elif isinstance(exc, (httpx.TimeoutException, asyncio.TimeoutError)):
                    message = "请求超时"
                elif isinstance(exc, httpx.TransportError):
                    message = "网络连接失败"
                else:
                    message = "来源格式变化或评分数据无效"
                entry.update(status="error", error=f"{message}，保留上次评分")
            self.db.save_benchmark_cache(source, entry)
            self.cache[source] = entry

        await asyncio.gather(
            sync("superclue", self._superclue),
            sync("bfcl", self._bfcl),
            sync("artificial_analysis", self._aa),
        )
